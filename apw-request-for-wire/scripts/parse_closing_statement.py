#!/usr/bin/env python3
"""
Parse APW Brasil Closing Statement Excel into structured JSON.

Usage:
    python3 parse_closing_statement.py <path_to_closing_statement.xlsx>

Output: JSON to stdout

Excel structure (1-indexed rows → 0-indexed in this script):
  Row 1  (idx 0)  : empty
  Row 2  (idx 1)  : "COMPLETION STATEMENT"
  Row 3  (idx 2)  : "AP WIRELESS BRASIL..."
  Row 4  (idx 3)  : "PURCHASE OF:" | "L1042188_LANDLORD" | ... | "Tenant" header
  Row 5  (idx 4)  : "CARRIER"     | ...                 | ATC  | months | rent | total
  Row 6  (idx 5)  : "PROPOSED COMPLETION DATE:" | | date
  Row 7  (idx 6)  : "LANDLORD:"
  Row 8  (idx 7)  : "PURCHASING ENTITY" | | | entity_name
  Row 9  (idx 8)  : "ITEMS:" header
  Row 10 (idx 9)  : currency labels (R$)
  Rows 11-19 (idx 10-18): line items (B=label, D=deduct, F=add)
  Row 20 (idx 19) : "(-) Installments"
  Row 21 (idx 20) : "Sub-total"
  Row 22 (idx 21) : "BALANCE DUE:" | | | total_balance
  Row 23 (idx 22) : empty
  Row 24+ (idx 23+): beneficiary blocks (repeating):
    "Total Amount due to Landlord at closing" → G = amount
    "Landlord Bank Details:"
    "Beneficiary:" → D = name
    "Bank:"        → D = bank_name
    "Bank Code:"   → D = code
    "Account:"     → D = account
    "Branch:"      → D = branch
    "Tax ID:"      → D = tax_id
"""

import sys
import json
import re
import openpyxl
from pathlib import Path


def to_float(value):
    """Convert cell value to float, handling None, strings, and formula strings."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        # Remove R$, spaces; convert BR format (dots=thousands, comma=decimal)
        cleaned = value.strip().replace('R$', '').replace(' ', '')
        cleaned = cleaned.replace('.', '').replace(',', '.')
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def parse_investment_sheet(ws):
    """Parse the Investment / Completion Statement sheet."""
    result = {
        "deal": {},
        "line_items": [],
        "pro_rations": [],
        "beneficiaries": [],
        "balance_due": None,
        "sub_total_deduct": None,
        "sub_total_add": None,
    }

    rows = list(ws.iter_rows(values_only=True))

    # ── Deal header (row 4, idx 3) ────────────────────────────────────────────
    # Col B = "PURCHASE OF:", Col C = "L1042188_JOSE ANDRE DE OLIVEIRA"
    deal_ref = str(rows[3][2] or "").strip()
    m = re.match(r'(L\d+)[_\s]+(.+)', deal_ref)
    if m:
        result["deal"]["l_number"] = m.group(1)
        result["deal"]["landlord"] = m.group(2).strip()
    else:
        result["deal"]["l_number"] = ""
        result["deal"]["landlord"] = deal_ref

    # ── Carrier / pro-rations (row 5, idx 4) ─────────────────────────────────
    # Col B = "CARRIER", Col J (idx 9) = "ATC", N(13)=months, O(14)=rent, P(15)=total
    carrier_row = rows[4]
    result["deal"]["carrier"] = str(carrier_row[9] or "").strip()

    months = to_float(carrier_row[13])
    rent_per_month = to_float(carrier_row[14])
    apportionment = to_float(carrier_row[15])

    if result["deal"]["carrier"] and result["deal"]["carrier"] not in ("Tenant", ""):
        result["pro_rations"].append({
            "tenant": result["deal"]["carrier"],
            "months": months,
            "rent_per_month": rent_per_month,
            "total_apportionment": apportionment,
        })

    # ── Completion date (row 6, idx 5) ───────────────────────────────────────
    completion_val = rows[5][3]  # Col D
    if completion_val:
        try:
            result["deal"]["completion_date"] = completion_val.strftime("%B %d, %Y")
        except AttributeError:
            result["deal"]["completion_date"] = str(completion_val)
    else:
        result["deal"]["completion_date"] = ""

    # ── Purchasing entity (row 8, idx 7) ─────────────────────────────────────
    result["deal"]["purchasing_entity"] = str(rows[7][3] or "").strip()  # Col D

    # ── Line items (rows 11-19, idx 10-18) ───────────────────────────────────
    # Col B = label, Col D (idx 3) = Deduct, Col F (idx 5) = Add
    LABEL_COL = 1   # B
    DEDUCT_COL = 3  # D
    ADD_COL = 5     # F
    SKIP_LABELS = {"items:", "sub-total", "balance due:", "(-) installments", ""}

    for i in range(10, 21):
        if i >= len(rows):
            break
        row = rows[i]
        label = str(row[LABEL_COL] or "").strip()
        if not label or label.lower() in SKIP_LABELS:
            continue

        deduct = to_float(row[DEDUCT_COL])
        add = to_float(row[ADD_COL])

        if deduct is None and add is None:
            continue

        net = (add or 0) + (deduct or 0)  # deduct is already negative in sheet

        result["line_items"].append({
            "label": label.strip(),
            "deduct": deduct,
            "add": add,
            "net": round(net, 2),
        })

    # ── Sub-total and Balance Due (rows 21-22, idx 20-21) ────────────────────
    if 20 < len(rows):
        st = rows[20]
        result["sub_total_deduct"] = to_float(st[DEDUCT_COL])
        result["sub_total_add"] = to_float(st[ADD_COL])

    if 21 < len(rows):
        result["balance_due"] = to_float(rows[21][DEDUCT_COL])  # Col D

    # ── Beneficiaries (row 24 onward, idx 23+) ───────────────────────────────
    current_bene = None
    beneficiaries = []

    for i in range(22, len(rows)):
        row = rows[i]
        label_raw = str(row[LABEL_COL] or "").strip()
        label = label_raw.lower()

        if "total amount due" in label:
            if current_bene and current_bene.get("name"):
                beneficiaries.append(current_bene)
            amount = to_float(row[6])  # Col G
            current_bene = {
                "amount": amount, "name": "", "bank": "", "bank_code": None,
                "branch": "", "account": "", "tax_id": ""
            }

        elif label == "beneficiary:" and current_bene is not None:
            current_bene["name"] = str(row[3] or "").strip()

        elif label == "bank:" and current_bene is not None:
            current_bene["bank"] = str(row[3] or "").strip()

        elif label == "bank code:" and current_bene is not None:
            current_bene["bank_code"] = str(row[3] or "").strip()

        elif label == "account:" and current_bene is not None:
            current_bene["account"] = str(row[3] or "").strip()

        elif label == "branch:" and current_bene is not None:
            current_bene["branch"] = str(row[3] or "").strip()

        elif label == "tax id:" and current_bene is not None:
            current_bene["tax_id"] = str(row[3] or "").strip()

    if current_bene and current_bene.get("name"):
        beneficiaries.append(current_bene)

    result["beneficiaries"] = beneficiaries

    # ── Reconcile beneficiary amounts if any are None ─────────────────────────
    known = sum(b["amount"] for b in beneficiaries if b.get("amount") is not None)
    for b in beneficiaries:
        if b["amount"] is None and result["balance_due"] is not None:
            others = sum(x["amount"] or 0 for x in beneficiaries if x is not b and x["amount"])
            b["amount"] = round(result["balance_due"] - others, 2)
            b["amount_inferred"] = True

    return result


def parse_checklist_sheet(ws):
    """Parse the Check List sheet — skip header rows."""
    checklist = []
    rows = list(ws.iter_rows(values_only=True))
    for row in rows[2:]:
        item_raw = str(row[0] or "").strip() if row[0] else ""
        if not item_raw or item_raw.startswith(" Request"):
            continue
        status = str(row[2] or "").strip()
        location = str(row[1] or "").strip()
        checklist.append({"item": item_raw, "status": status, "location": location})
    return checklist


def parse_closing_statement(filepath):
    """Main entry: parse all relevant sheets."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    output = {}

    # Find and parse the Investment sheet
    investment_sheet = None
    for name in wb.sheetnames:
        if name.lower() in ("investment", "completion statement", "closing statement"):
            investment_sheet = wb[name]
            break
    if investment_sheet is None:
        investment_sheet = wb.worksheets[0]

    output.update(parse_investment_sheet(investment_sheet))

    # Find and parse the Check List sheet
    for name in wb.sheetnames:
        if "check" in name.lower():
            output["checklist"] = parse_checklist_sheet(wb[name])
            break
    if "checklist" not in output:
        output["checklist"] = []

    # Infer deal type from filename
    filename = Path(filepath).stem
    fn_lower = filename.lower()
    if "surface right" in fn_lower or "drs" in fn_lower:
        output["deal"]["deal_type"] = "Surface Right Deed"
    elif "assignment" in fn_lower or "cdc" in fn_lower:
        output["deal"]["deal_type"] = "Assignment of Rents Agreement"
    elif "fee simple" in fn_lower or "compra" in fn_lower:
        output["deal"]["deal_type"] = "Fee Simple / Purchase Agreement"
    else:
        output["deal"]["deal_type"] = "Purchase Agreement"

    # Fallback: infer L-number from filename if not in sheet
    if not output["deal"].get("l_number"):
        m = re.search(r'L\d{5,}', filename)
        if m:
            output["deal"]["l_number"] = m.group(0)

    # Fallback: infer landlord from filename if not in sheet
    if not output["deal"].get("landlord"):
        # Filename format: "Closing Statement - L1042188 - JOSE ANDRE - ..."
        parts = filename.split(" - ")
        if len(parts) >= 3:
            output["deal"]["landlord"] = parts[2].strip()

    return output


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python3 parse_closing_statement.py <path>", file=sys.stderr)
        sys.exit(1)

    data = parse_closing_statement(sys.argv[1])
    print(json.dumps(data, indent=2, ensure_ascii=False))
