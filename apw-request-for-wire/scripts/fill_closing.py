#!/usr/bin/env python3
"""
fill_closing.py — gera o Closing Statement (padrão Nilton) + rascunho do e-mail RFW
a partir de um JSON do deal.

Uso:
  python3 fill_closing.py deal.json --out <pasta> [--template references/closing_template.xlsx]

Saídas (na pasta --out):
  Closing Statement - <L> - <Nome> - <Tipo>.xlsx
  RFW - <L> - <Nome>.html          (corpo do e-mail, pronto p/ colar no Outlook)
  stdout: resumo de validação (Balance Due x blocos, custas, datas)

O template é um closing real do Nilton com as células variáveis limpas. Linhas 2-19
nunca mudam de posição; a seção de pagamento (linha 20 em diante) é reconstruída
copiando os estilos do próprio template. Qualquer closing antigo do Nilton também
serve como --template (as células variáveis são sobrescritas).

Schema do JSON: ver references/deal_schema.md
"""
import argparse, copy, datetime as dt, html, json, os, re, sys
import openpyxl
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_TEMPLATE = os.path.join(HERE, "..", "references", "closing_template.xlsx")

MONTHS_EN = ["January", "February", "March", "April", "May", "June", "July",
             "August", "September", "October", "November", "December"]

# ---------------------------------------------------------------- helpers
def fnum(v):
    """número -> texto p/ fórmula sem perder centavos (12195.51, não 12195.5)."""
    return format(float(v), ".2f").rstrip("0").rstrip(".")

def brl(v):
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if v < 0 else "") + s

def num_or_formula(parts):
    """lista de componentes -> '=a+b+c' (padrão Nilton) ou número."""
    if parts is None:
        return None, 0.0
    if isinstance(parts, (int, float)):
        return parts, float(parts)
    parts = [p for p in parts if p not in (None, 0)]
    if not parts:
        return 0, 0.0
    if len(parts) == 1:
        return parts[0], float(parts[0])
    return "=" + "+".join(fnum(p) for p in parts), float(sum(parts))

def snap_styles(ws, row):
    return {c: copy.copy(ws.cell(row=row, column=c)._style) for c in range(2, 8)}

def apply(ws, row, styles, height=None):
    for c, st in styles.items():
        ws.cell(row=row, column=c)._style = copy.copy(st)
    if height:
        ws.row_dimensions[row].height = height

# ---------------------------------------------------------------- core
def build(spec, template, outdir):
    wb = openpyxl.load_workbook(template)
    ws = wb["Investment"]

    # estilos-modelo da seção inferior (capturados ANTES de limpar)
    st_item, st_sub, st_bal = snap_styles(ws, 20), snap_styles(ws, 21), snap_styles(ws, 22)
    st_amt, st_hdr = snap_styles(ws, 24), snap_styles(ws, 25)
    st_field_first, st_field, st_field_last = snap_styles(ws, 26), snap_styles(ws, 27), snap_styles(ws, 31)

    # limpa seção inferior (linha 20+) — valores, estilos e merges
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 20:
            ws.unmerge_cells(str(rng))
    for r in range(20, max(ws.max_row, 60) + 1):
        for c in range(2, 8):
            cell = ws.cell(row=r, column=c)
            cell.value = None
            cell._style = copy.copy(ws.cell(row=200, column=c)._style)
        ws.row_dimensions[r].height = None

    L, name, dtype = spec["l_code"], spec["name"], spec["deal_type"]
    ws["C4"] = f"{L}_{name}"
    d6 = dt.date.fromisoformat(spec["completion_date"])
    ws["D6"] = dt.datetime(d6.year, d6.month, d6.day)

    # painel pró-rata J5:P5
    pr = spec.get("prorata")
    ws["J5"] = spec["tenant"]
    if pr:
        ws["K5"], ws["L5"], ws["M5"] = pr["from"], pr["to"], pr["completion"]
        ws["N5"], ws["O5"] = pr["months"], pr["rent"]
        prorata_total = round(pr["months"] * pr["rent"], 2)
        months_lbl = pr.get("label") or _month_span(pr["from"], pr["months"])
        ws["B12"] = "Rent Apportionments" + " " * 50 + months_lbl
        ws["D12"] = "=-P5"
    else:
        ws["K5"] = ws["L5"] = ws["M5"] = "-"
        ws["N5"] = ws["O5"] = 0
        prorata_total = 0.0
        ws["B12"] = "Rent Apportionments"
        ws["D12"] = 0

    # preço
    pv, price = num_or_formula(spec["price"])
    ws["F11"] = pv

    # linha 13: Broker da APW (CAPEX TBSA descontinuado)
    r13 = spec.get("row13")
    wire_extra = 0.0          # valores em F sem D (entram no Balance Due)
    if r13:
        ws["B13"] = "Intermediary / Broker Fee"
        ws["F13"] = r13["amount"]
        if r13.get("paid_in_wire"):      # broker APW pago no próprio wire (UW: sem D)
            ws["D13"] = None
            wire_extra += r13["amount"]
        else:
            ws["D13"] = "=-F13"
    else:
        ws["B13"] = "Intermediary / Broker Fee"
        ws["D13"] = ws["F13"] = None

    # custas 14-19 (D = -F), auxílio 18 sem D
    costs = spec["costs"]
    tot_costs = 0.0
    for row, key in [(14, "title"), (15, "notary"), (16, "recording"), (17, "legal"), (19, "site_inspection")]:
        v, f = num_or_formula(costs.get(key, 0))
        ws[f"F{row}"] = v
        ws[f"D{row}"] = f"=-F{row}"
        tot_costs += f
    aux = spec.get("legal_aid", {}) or {}
    aux_amt = float(aux.get("amount", 0) or 0)
    ws["F18"] = aux_amt
    ws["D18"] = None

    # ---------------- seção inferior
    r = 20
    adv_rows = []
    for adv in spec.get("advance_payments", []):          # sinal já pago (Antônio)
        apply(ws, r, st_item)
        ws[f"B{r}"] = adv["label"]
        ws[f"D{r}"] = -abs(adv["amount"])
        adv_rows.append(r)
        r += 1
    inst_row = r
    apply(ws, r, st_item); ws[f"B{r}"] = "(-) Installments"
    r += 1
    sub_row = r
    apply(ws, r, st_sub); ws[f"B{r}"] = "Sub-total"
    ws[f"D{r}"] = f"=SUM(D11:D{r-1})"; ws[f"F{r}"] = f"=SUM(F11:F{r-1})"
    ws[f"E{r}"] = ws[f"G{r}"] = " "
    r += 1
    bal_row = r
    apply(ws, r, st_bal, 15); ws[f"B{r}"] = "BALANCE DUE:"
    ws[f"D{r}"] = f"=SUM(D{sub_row},F{sub_row})"
    ws.merge_cells(f"D{r}:G{r}")
    r += 2

    # ---- bloco do LL
    first = spec.get("first_payment")          # None = preço cheio
    attached = spec.get("ll_attached_lines", [])   # IPTU, condicionado, material técnico (sem banco)
    installments = spec.get("installments", [])
    others = spec.get("other_payees", [])          # advogado, broker do LL, broker APW-wire...

    ll_row = r
    apply(ws, r, st_amt); ws.merge_cells(f"C{r}:F{r}")
    ws[f"B{r}"] = "Total Amount due to Landlord at closing"
    r += 1
    att_rows = []
    for a in attached:
        apply(ws, r, st_amt); ws.merge_cells(f"C{r}:F{r}")
        ws[f"B{r}"] = a["label"]; ws[f"G{r}"] = a["amount"]
        att_rows.append(r); r += 1
    inst_rows = []
    for it in installments:
        apply(ws, r, st_amt); ws.merge_cells(f"C{r}:F{r}")
        ws[f"B{r}"] = it["label"]; ws[f"G{r}"] = it["amount"]
        inst_rows.append(r); r += 1
    if inst_rows:
        ws[f"D{inst_row}"] = f"=-SUM(G{inst_rows[0]}:G{inst_rows[-1]})"
    banks = spec["landlord_bank"] if isinstance(spec["landlord_bank"], list) else [spec["landlord_bank"]]
    for b in banks:
        r = _bank_block(ws, r, b, st_hdr, st_field_first, st_field, st_field_last)

    other_rows = []
    for o in others:
        apply(ws, r, st_amt); ws.merge_cells(f"C{r}:F{r}")
        ws[f"B{r}"] = o.get("label", "Total Amount due to Landlord at closing")
        if o.get("c_label"):
            ws[f"C{r}"] = o["c_label"]
        ws[f"G{r}"] = o["amount"]
        other_rows.append((r, o)); r += 1
        for x in o.get("extra_lines", []):          # ex.: 2ª metade do broker condicionada
            apply(ws, r, st_amt); ws.merge_cells(f"C{r}:F{r}")
            ws[f"B{r}"] = x["label"]; ws[f"G{r}"] = x["amount"]
            other_rows.append((r, x)); r += 1
        if o.get("bank"):
            r = _bank_block(ws, r, o["bank"], st_hdr, st_field_first, st_field, st_field_last,
                            header=o.get("bank_header", "Bank Details:"))

    # fórmula do LL (padrão Nilton)
    base = fnum(first) if first is not None else "F11"
    f = f"={base}+D12+F18"
    for ar in adv_rows:
        f += f"+D{ar}"
    for ar, a in zip(att_rows, attached):
        if a.get("carve_from_landlord", True):
            f += f"-G{ar}"
    for orow, o in other_rows:
        if o.get("carve_from_landlord", True):
            f += f"-G{orow}"
    ws[f"G{ll_row}"] = f

    # ---------------- cálculo espelho (validação)
    adv_total = sum(abs(a["amount"]) for a in spec.get("advance_payments", []))
    inst_total = sum(i["amount"] for i in installments)
    balance = price - prorata_total + aux_amt + wire_extra - adv_total - inst_total
    ll = (first if first is not None else price) - prorata_total + aux_amt - adv_total \
         - sum(a["amount"] for a in attached if a.get("carve_from_landlord", True)) \
         - sum(o["amount"] for _, o in other_rows if o.get("carve_from_landlord", True))
    blocks = ll + sum(a["amount"] for a in attached) + sum(o["amount"] for _, o in other_rows)

    # ---------------- check list
    cl = wb["Check List"]
    for k, v in (spec.get("checklist") or {}).items():
        for rr in range(3, cl.max_row + 1):
            if cl[f"A{rr}"].value and k.lower() in str(cl[f"A{rr}"].value).lower():
                cl[f"C{rr}"] = v

    wb.calculation.fullCalcOnLoad = True
    os.makedirs(outdir, exist_ok=True)
    fname = f"Closing Statement - {L} - {name} - {dtype}.xlsx"
    path = os.path.join(outdir, fname)
    wb.save(path)

    report = {
        "file": path, "balance_due": round(balance, 2), "landlord_at_closing": round(ll, 2),
        "sum_of_blocks": round(blocks, 2), "closing_costs_net_zero": round(tot_costs, 2),
        "legal_aid": aux_amt, "prorata": prorata_total, "installments_future": round(inst_total, 2),
    }
    warnings = _validate(spec, report, installments, d6)
    email_path = write_email(spec, report, outdir, ws_rows=dict(ll_row=ll_row))
    report["email"] = email_path
    return report, warnings

def _month_span(first, n):
    i = MONTHS_EN.index(first)
    ms = [MONTHS_EN[(i + k) % 12] for k in range(n)]
    return ", ".join(ms)

def _bank_block(ws, r, b, st_hdr, st_first, st_field, st_last, header="Landlord Bank Details:"):
    apply(ws, r, st_hdr); ws.merge_cells(f"D{r}:G{r}")
    ws[f"B{r}"] = b.get("header", header)
    if b.get("note"):
        ws[f"D{r}"] = b["note"]                    # ex.: "50% of each payment"
    r += 1
    fields = [("Beneficiary:", "beneficiary"), ("Bank:", "bank"), ("Bank Code:", "bank_code"),
              ("Account:", "account"), ("Branch:", "branch"), ("Tax ID:", "tax_id")]
    for i, (lbl, k) in enumerate(fields):
        st = st_first if i == 0 else (st_last if i == len(fields) - 1 else st_field)
        apply(ws, r, st); ws.merge_cells(f"D{r}:G{r}")
        ws[f"B{r}"] = lbl
        ws[f"D{r}"] = b.get(k)
        r += 1
    return r

def _validate(spec, rep, installments, d6):
    w = []
    if abs(rep["balance_due"] - rep["sum_of_blocks"]) > 0.01:
        w.append(f"🔴 Soma dos blocos ({brl(rep['sum_of_blocks'])}) ≠ Balance Due ({brl(rep['balance_due'])})")
    if rep["landlord_at_closing"] <= 0:
        w.append("🔴 Valor ao LL ≤ 0")
    ci = spec.get("custas_items")
    if ci is not None:
        exp = round(sum(ci), 2)
        if abs(exp - rep["closing_costs_net_zero"]) > 0.01:
            w.append(f"🔴 Custas linhas 14-17+19 = {brl(rep['closing_costs_net_zero'])} ≠ itens da tabela do jurídico (sem auxílio) = {brl(exp)}")
    if spec.get("custas_saldo_juridico") is not None and ci is not None:
        if abs(spec["custas_saldo_juridico"] - (sum(ci) + rep["legal_aid"])) > 0.01:
            w.append(f"🟡 'Saldo' do e-mail do jurídico ({brl(spec['custas_saldo_juridico'])}) não fecha com a soma dos itens — usei a soma dos itens")
    today = dt.date.today()
    if abs((d6 - today).days) > 120:
        w.append(f"🟡 Proposed completion date {d6} está a >120 dias de hoje — confirmar")
    years = [int(m) for it in installments for m in re.findall(r"/(\d{4})", it["label"])]
    if years and (years != sorted(years) or max(years) - min(years) > 40):
        w.append(f"🔴 Anos das parcelas fora de ordem/suspeitos: {years}")
    if years and any(y < d6.year for y in years):
        w.append("🔴 Parcela com ano anterior ao closing")
    return w

# ---------------------------------------------------------------- e-mail
def write_email(spec, rep, outdir, ws_rows=None):
    """Corpo HTML no formato do e-mail que o Nilton envia (Request for Wire)."""
    L, name, dtype = spec["l_code"], spec["name"], spec["deal_type"]
    fd = dt.date.fromisoformat(spec.get("funding_date") or spec["completion_date"])
    fd_txt = f"{fd.strftime('%A')}, {MONTHS_EN[fd.month-1]}/{fd.day:02d}/{fd.year}"
    sign_word = {"Assignment of Rents": "Assignment of Rents",
                 "Surface Right": "Surface Right",
                 "Fee Simple Interest": "Fee Simple Interest",
                 "Credit Assignment and Surface Right": "Credit Assignment and Surface Right"}.get(dtype, dtype)
    e = html.escape
    notes = spec.get("please_note") or []
    rows_items = _items_rows(spec, rep)
    pay_html = _payment_html(spec, rep)
    notes_html = ""
    if notes:
        notes_html = "<p><b>Please note:</b></p><ul style='margin-top:0'>" + \
            "".join(f"<li><b>{e(n)}</b></li>" for n in notes) + "</ul>"
    body = f"""<html><body style="font-family:Calibri,Arial,sans-serif;font-size:11pt">
<p>Team,</p>
<p>We are about to sign <b>{e(sign_word)}</b> pertaining to this transaction.</p>
<p>Expected Funding Date <span style="background:#ffff00;color:#c00000;font-weight:bold;text-decoration:underline;font-size:13pt">{e(fd_txt)}</span></p>
<p>Final Closing Costs:</p>
{rows_items}
<p><b>{e(spec.get('key_notes',''))}</b></p>
{notes_html}
<p style="color:#c00000"><b>Once approved, funds should be directed to the following bank accounts:</b></p>
{pay_html}
<p>If you have any questions, please let us know.</p>
</body></html>"""
    p = os.path.join(outdir, f"RFW - {L} - {name}.html")
    with open(p, "w", encoding="utf-8") as fh:
        fh.write(body)
    # subject / destinatários em txt ao lado
    with open(os.path.join(outdir, f"RFW - {L} - {name} - header.txt"), "w", encoding="utf-8") as fh:
        fh.write(f"Subject: Request for Wire - {L} - {name} - {dtype}\n")
        fh.write("To: " + "; ".join(RFW_TO) + "\nCc: " + "; ".join(RFW_CC) + "\n")
    return p

# Distribuição padrão (nomes; o Outlook resolve os endereços — repo é público, não versionar e-mails)
RFW_TO = ["Paulina Guerrero", "Juan Carlos Martinez", "Ryan Eppel", "Luiz Alfredo", "Fabiana Siviero"]
RFW_CC = ["Scott Langeland", "Dan Hasselman", "Esteban Castrejon", "Eric Pape", "Michelle Carvalho",
          "Marcelina Silva", "Lucas Minozzo", "Daniel Presson", "Guglielmo Noya",
          "Marnique Richards-Ouellette", "Diogo Bueno", "Douglas Alves", "Gracia Serrano", "Daniel Romagnolo"]

TD = "border:1px solid #000;padding:2px 6px;vertical-align:bottom"

def _items_rows(spec, rep):
    e = html.escape
    c = spec["costs"]
    def tot(k):
        return num_or_formula(c.get(k, 0))[1]
    pr = rep["prorata"]
    pr_lbl = spec["prorata"].get("label") or _month_span(spec["prorata"]["from"], spec["prorata"]["months"]) if spec.get("prorata") else ""
    d6 = dt.date.fromisoformat(spec["completion_date"])
    d6_pt = _date_pt(d6)
    r13 = spec.get("row13")
    lines = [
        ("Purchase Price / Valor Investimento", "Lease M-files", "", "R$ " + brl(num_or_formula(spec["price"])[1])),
        ("Rent Apportionments " + pr_lbl, "Lease M-files", f"<span style='color:#c00000'>{brl(-pr)}</span>" if pr else "0,00", ""),
        ("Intermediary / Broker Fee", "",
         (brl(-r13['amount']) if r13 and not r13.get('paid_in_wire') else ""), (brl(r13['amount']) if r13 else "")),
        ("Title Search Charges/Certidões/Documentos/Correio", "link", brl(-tot("title")), brl(tot("title"))),
        ("Local Notary Fees/ Despesas com Cartório (s)", "Link", brl(-tot("notary")), brl(tot("notary"))),
        ("Recording Fees/ Custo de Registro", "Link", brl(-tot("recording")), brl(tot("recording"))),
        ("Legal Fees", "Fixed Amount", brl(-tot("legal")), brl(tot("legal"))),
        ("Site Owner Legal Fee/ Auxílio Advogado", "Director", "", brl(rep["legal_aid"])),
        ("Site Inspection", "Link", brl(-tot("site_inspection")), brl(tot("site_inspection"))),
    ]
    for a in spec.get("advance_payments", []):
        lines.append((a["label"], "", brl(-abs(a["amount"])), ""))
    lines.append(("(-) Installments", "", brl(-rep["installments_future"]) if rep["installments_future"] else "", ""))
    rows = "".join(f"<tr><td style='{TD}'>{e(a)}</td><td style='{TD};text-align:center'>{e(b)}</td>"
                   f"<td style='{TD};text-align:right'>{cc}</td><td style='{TD}'></td>"
                   f"<td style='{TD};text-align:right'>{d}</td><td style='{TD}'></td></tr>" for a, b, cc, d in lines)
    return f"""<table style="border-collapse:collapse;font-size:10pt;min-width:820px">
<tr><td colspan=6 style="{TD};text-align:center;font-size:13pt;font-weight:bold;height:50px;vertical-align:bottom">COMPLETION STATEMENT</td></tr>
<tr><td colspan=6 style="{TD};text-align:center;font-weight:bold">AP WIRELESS BRASIL INVESTIMENTOS IMOBILIÁRIOS LTDA.</td></tr>
<tr><td style="{TD};font-weight:bold">PURCHASE OF:</td><td colspan=5 style="{TD};text-align:center;font-weight:bold;font-size:12pt">{e(spec['l_code'])}_{e(spec['name'])}</td></tr>
<tr><td style="{TD};font-weight:bold">CARRIER</td><td style="{TD}"></td><td colspan=4 style="{TD}"></td></tr>
<tr><td style="{TD};font-weight:bold">PROPOSED COMPLETION DATE:</td><td style="{TD}"></td><td colspan=4 style="{TD};color:#c00000">{e(d6_pt)}</td></tr>
<tr><td style="{TD};font-weight:bold">LANDLORD:</td><td style="{TD}"></td><td colspan=4 style="{TD}"></td></tr>
<tr><td style="{TD};font-weight:bold">PURCHASING ENTITY</td><td style="{TD}"></td><td colspan=4 style="{TD};text-align:center">AP Wireless Brasil Investimentos Imobiliario</td></tr>
<tr><td style="{TD};font-weight:bold">ITEMS:</td><td style="{TD};font-weight:bold;text-align:center">Source/Fonte</td><td style="{TD};text-align:center">Deduct</td><td style="{TD};text-align:center">Deduct (VAT)</td><td style="{TD};text-align:center">Add</td><td style="{TD};text-align:center">ADD (VAT)</td></tr>
<tr><td style="{TD}"></td><td style="{TD}"></td><td style="{TD};text-align:center">R$</td><td style="{TD};text-align:center">R$</td><td style="{TD};text-align:center">R$</td><td style="{TD};text-align:center">R$</td></tr>
{rows}
<tr><td style="{TD};text-align:right">Sub-total</td><td style="{TD}"></td><td style="{TD};text-align:right">{brl(-(rep['prorata']+rep['closing_costs_net_zero']+rep['installments_future']+sum(abs(a['amount']) for a in spec.get('advance_payments',[]))+((spec['row13']['amount']) if spec.get('row13') and not spec['row13'].get('paid_in_wire') else 0)))}</td><td style="{TD}"></td><td style="{TD};text-align:right">{brl(num_or_formula(spec['price'])[1]+rep['closing_costs_net_zero']+rep['legal_aid']+((spec['row13']['amount']) if spec.get('row13') else 0))}</td><td style="{TD}"></td></tr>
<tr><td style="{TD};text-align:right;font-weight:bold">BALANCE DUE:</td><td style="{TD}"></td><td colspan=4 style="{TD};text-align:right;font-weight:bold">R$ {brl(rep['balance_due'])}</td></tr>
</table>"""

def _payment_html(spec, rep):
    e = html.escape
    red = "color:#c00000"
    out = []
    def amt_row(label, amount, c_label=""):
        return (f"<tr><td style='{TD};{red};font-weight:bold;text-decoration:underline'>{e(label)}</td>"
                f"<td colspan=2 style='{TD}'>{e(c_label)}</td><td style='{TD};{red};font-weight:bold;text-align:right'>R$ {brl(amount)}</td></tr>")
    def bank(b, header="Landlord Bank Details:"):
        s = f"<tr><td style='{TD};{red}'>{e(b.get('header', header))}</td><td style='{TD}'></td><td colspan=2 style='{TD}'>{e(b.get('note',''))}</td></tr>"
        for lbl, k in [("Beneficiary:", "beneficiary"), ("Bank:", "bank"), ("Bank Code:", "bank_code"),
                       ("Account:", "account"), ("Branch:", "branch"), ("Tax ID:", "tax_id")]:
            bold = "font-weight:bold" if k == "beneficiary" else ""
            s += f"<tr><td style='{TD};{red}'>{lbl}</td><td style='{TD}'></td><td colspan=2 style='{TD};{bold}'>{e(str(b.get(k,'')))}</td></tr>"
        return s
    out.append(amt_row("Total Amount due to Landlord at closing", rep["landlord_at_closing"]))
    for a in spec.get("ll_attached_lines", []):
        out.append(amt_row(a["label"], a["amount"]))
    for it in spec.get("installments", []):
        out.append(amt_row(it["label"], it["amount"]))
    banks = spec["landlord_bank"] if isinstance(spec["landlord_bank"], list) else [spec["landlord_bank"]]
    for b in banks:
        out.append(bank(b))
    for o in spec.get("other_payees", []):
        out.append(amt_row(o.get("label", "Total Amount due at closing"), o["amount"], o.get("c_label", "")))
        for x in o.get("extra_lines", []):
            out.append(amt_row(x["label"], x["amount"]))
        if o.get("bank"):
            out.append(bank(o["bank"], o.get("bank_header", "Bank Details:")))
    return "<table style='border-collapse:collapse;font-size:10pt;min-width:820px'>" + "".join(out) + "</table>"

def _date_pt(d):
    dias = ["segunda-feira", "terça-feira", "quarta-feira", "quinta-feira", "sexta-feira", "sábado", "domingo"]
    meses = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto",
             "setembro", "outubro", "novembro", "dezembro"]
    return f"{dias[d.weekday()]}, {d.day} de {meses[d.month-1]} de {d.year}"

# ---------------------------------------------------------------- main
if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("--out", required=True)
    ap.add_argument("--template", default=DEFAULT_TEMPLATE)
    a = ap.parse_args()
    with open(a.spec, encoding="utf-8") as fh:
        spec = json.load(fh)
    rep, warns = build(spec, a.template, a.out)
    print(json.dumps(rep, ensure_ascii=False, indent=2))
    print("\nVALIDAÇÃO:")
    print("\n".join(warns) if warns else "✅ tudo fecha (blocos = Balance Due; custas ok; datas ok)")
    sys.exit(1 if any(w.startswith("🔴") for w in warns) else 0)
